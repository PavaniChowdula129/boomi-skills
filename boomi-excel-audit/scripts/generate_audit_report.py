import argparse
import sys
import os
import json
import base64
import urllib.request
import urllib.error
import xml.etree.ElementTree as ET
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

def get_auth_header(username: str, token: str) -> str:
    raw = f"BOOMI_TOKEN.{username}:{token}"
    encoded = base64.b64encode(raw.encode("utf-8")).decode("ascii")
    return f"Basic {encoded}"

def call_boomi_api(endpoint: str, payload: dict, account_id: str, base_url: str, username: str, token: str) -> dict:
    url = f"{base_url.rstrip('/')}/api/rest/v1/{account_id}/{endpoint}"
    req = urllib.request.Request(
        url,
        data=json.dumps(payload).encode("utf-8") if payload is not None else None,
        headers={
            "Authorization": get_auth_header(username, token),
            "Content-Type": "application/json",
            "Accept": "application/json"
        },
        method="POST" if payload is not None else "GET"
    )
    try:
        with urllib.request.urlopen(req) as resp:
            return json.loads(resp.read().decode("utf-8"))
    except urllib.error.HTTPError as e:
        sys.stderr.write(f"API Call failed [{e.code}] for {endpoint}: {e.read().decode('utf-8')}\n")
        return {}
    except Exception as e:
        sys.stderr.write(f"Connection error calling {endpoint}: {str(e)}\n")
        return {}

def get_all_environments(account_id: str, base_url: str, username: str, token: str) -> list:
    query = {"QueryFilter": {}}
    res = call_boomi_api("Environment/query", query, account_id, base_url, username, token)
    envs = []
    for env in res.get("result", []):
        envs.append({
            "id": env.get("id"),
            "name": env.get("name", "Unknown Environment")
        })
    return envs

def fetch_component_xml(comp_id: str, account_id: str, base_url: str, username: str, token: str) -> str:
    url = f"{base_url.rstrip('/')}/api/rest/v1/{account_id}/Component/{comp_id}"
    req = urllib.request.Request(
        url,
        headers={
            "Authorization": get_auth_header(username, token),
            "Accept": "application/xml"
        },
        method="GET"
    )
    try:
        with urllib.request.urlopen(req) as resp:
            return resp.read().decode("utf-8")
    except Exception as e:
        sys.stderr.write(f"Failed to fetch XML for {comp_id}: {str(e)}\n")
        return ""

def find_references_in_xml(xml_str: str, child_comp_id: str, required_route_key: str = None) -> list:
    refs = []
    if not xml_str or child_comp_id not in xml_str:
        return refs
        
    try:
        root = ET.fromstring(xml_str)
        parent_map = {c: p for p in root.iter() for c in p}
        
        def get_context(elem):
            curr = elem
            details = []
            shape_context = None
            
            while curr is not None:
                tag = curr.tag
                if tag == "shape":
                    stype = curr.get("shapetype", "unknown")
                    slabel = curr.get("userlabel") or curr.get("name", "Unnamed")
                    shape_context = f"[{stype}] {slabel}"
                elif tag == "parametervalue":
                    name = curr.get("elementToSetName") or curr.get("key") or "Unnamed"
                    details.insert(0, f"Parameter: {name}")
                elif tag == "processparameter":
                    name = curr.get("processproperty", "Unknown")
                    details.insert(0, f"ProcessProperty: {name}")
                elif tag == "profileelement":
                    name = curr.get("elementName", "Unknown")
                    details.insert(0, f"ProfileElement: {name}")
                elif tag == "processProperty":
                    name = curr.get("name", "Unknown Property")
                    details.insert(0, f"[Process Property] {name}")
                elif tag == "dynamicProperty":
                    name = curr.get("name", "Unknown Property")
                    details.insert(0, f"[Dynamic Property] {name}")
                elif tag == "processRoute":
                    details.insert(0, f"[Process Route] {curr.get('name', 'Unnamed')}")
                elif tag == "BusinessRule":
                    details.insert(0, f"[Business Rule] {curr.get('name', 'Unnamed')}")
                elif tag == "mapFunction":
                    details.insert(0, f"MapFunction: {curr.get('name', 'Unnamed')}")
                elif tag == "Component":
                    ctype = curr.get("type", "unknown")
                    cname = curr.get("name", "Unnamed")
                    if not shape_context:
                        shape_context = f"[{ctype}] {cname}"
                
                curr = parent_map.get(curr)
                
            if shape_context:
                if details:
                    return f"{shape_context} -> " + " -> ".join(details)
                return shape_context
            elif details:
                return " -> ".join(details)
            return "(Configuration)"

        found_contexts = set()
        for elem in root.iter():
            is_match = False
            for k, v in elem.attrib.items():
                if v and child_comp_id in v:
                    is_match = True
                    break
            if not is_match and elem.text and child_comp_id in elem.text:
                is_match = True
                
            if is_match:
                if required_route_key is not None and elem.tag == 'processroutecall':
                    param = elem.find('.//routeParameter//staticparameter')
                    if param is not None:
                        shape_route_key = param.get('staticproperty')
                        if shape_route_key != required_route_key:
                            continue
                found_contexts.add(get_context(elem))
                
        valid_contexts = [c for c in found_contexts if c != "(Configuration)"]
        refs = valid_contexts if valid_contexts else (list(found_contexts) if found_contexts else [])
    except Exception as e:
        sys.stderr.write(f"Error parsing XML for shapes: {str(e)}\n")
        
    return refs

def resolve_component_metadata(comp_id: str, account_id: str, base_url: str, username: str, token: str) -> dict:
    query = {
        "QueryFilter": {
            "expression": {
                "operator": "EQUALS",
                "property": "componentId",
                "argument": [comp_id]
            }
        }
    }
    res = call_boomi_api("ComponentMetadata/query", query, account_id, base_url, username, token)
    records = res.get("result", [])
    if records:
        return records[0]
    return {"componentId": comp_id, "name": "Unknown", "type": "Unknown", "folderName": "Unknown"}

def get_deployed_packages(env_id: str, account_id: str, base_url: str, username: str, token: str) -> dict:
    if not env_id:
        return {}
    query = {
        "QueryFilter": {
            "expression": {
                "operator": "and",
                "nestedExpression": [
                    {"operator": "EQUALS", "property": "environmentId", "argument": [env_id]},
                    {"operator": "EQUALS", "property": "active", "argument": ["true"]}
                ]
            }
        }
    }
    res = call_boomi_api("DeployedPackage/query", query, account_id, base_url, username, token)
    deployed = {}
    for pkg in res.get("result", []):
        cid = pkg.get("componentId")
        deployed[cid] = {
            "version": pkg.get("componentVersion"),
            "packageId": pkg.get("packageId"),
            "packageVersion": pkg.get("packageVersion"),
            "name": pkg.get("componentName")
        }
    return deployed

class Node:
    def __init__(self, comp_id, name, type_, folder):
        self.comp_id = comp_id
        self.name = name
        self.type = type_
        self.folder = folder
        self.parents = [] # List of tuples: (parent_node, ref_type)
        self.is_circular = False

def build_graph(target_id, account_id, base_url, username, token):
    target_meta = resolve_component_metadata(target_id, account_id, base_url, username, token)
    root_node = Node(target_id, target_meta.get("name", target_id), target_meta.get("type", "Unknown"), target_meta.get("folderName", "Unknown"))
    
    node_cache = {target_id: root_node}
    
    def traverse(current_id, current_node, visited_path):
        if current_id in visited_path:
            current_node.is_circular = True
            return
        
        visited_path.add(current_id)
        
        query = {
            "QueryFilter": {
                "expression": {
                    "operator": "EQUALS",
                    "property": "componentId",
                    "argument": [current_id]
                }
            }
        }
        ref_res = call_boomi_api("ComponentReference/query", query, account_id, base_url, username, token)
        refs = ref_res.get("result", [])
        
        parent_map = {}
        for ref_container in refs:
            for ref in ref_container.get("references", []):
                parent_id = ref.get("parentComponentId")
                ref_type = ref.get("type", "DEPENDENT")
                if parent_id not in parent_map:
                    parent_map[parent_id] = set()
                parent_map[parent_id].add(ref_type)
                
        for parent_id, ref_types in parent_map.items():
            p_meta = resolve_component_metadata(parent_id, account_id, base_url, username, token)
            parent_type = p_meta.get("type", "Unknown")
            
            route_key = None
            cache_key = parent_id
            
            if parent_type == "processroute":
                p_xml = fetch_component_xml(parent_id, account_id, base_url, username, token)
                if p_xml:
                    try:
                        import xml.etree.ElementTree as ET
                        xml_body = p_xml[p_xml.find('<bns:object>')+12:p_xml.find('</bns:object>')] if '<bns:object>' in p_xml else p_xml
                        root_pr = ET.fromstring(xml_body)
                        for entry in root_pr.findall('.//processEntry'):
                            if entry.get('processId') == current_id:
                                route_key = entry.get('key')
                                break
                    except Exception:
                        pass
                if route_key:
                    cache_key = f"{parent_id}::route_{route_key}"

            if cache_key not in node_cache:
                p_node = Node(parent_id, p_meta.get("name", parent_id), parent_type, p_meta.get("folderName", "Unknown"))
                if route_key:
                    p_node.route_key = route_key
                node_cache[cache_key] = p_node
                traverse(parent_id, p_node, visited_path.copy())
            else:
                p_node = node_cache[cache_key]
                
            if parent_id in visited_path:
                p_node.is_circular = True
            
            added_edges = False
            p_xml = fetch_component_xml(parent_id, account_id, base_url, username, token)
            required_route_key = getattr(current_node, 'route_key', None)
            shapes = find_references_in_xml(p_xml, current_id, required_route_key=required_route_key)
            for shape_str in shapes:
                edge = (p_node, f"SHAPE:{shape_str}")
                if edge not in current_node.parents:
                    current_node.parents.append(edge)
                    added_edges = True
            
            if not added_edges:
                for r_type in ref_types:
                    edge = (p_node, r_type)
                    if edge not in current_node.parents:
                        current_node.parents.append(edge)
    
    traverse(target_id, root_node, set())
    return root_node, node_cache

def extract_paths(node, current_path, all_paths):
    if node.is_circular:
        all_paths.append(current_path + [{"node": node, "ref_type": "CIRCULAR_LINK"}])
        return
        
    if not node.parents:
        all_paths.append(current_path)
        return
        
    for parent_node, ref_type in node.parents:
        new_path = current_path + [{"node": parent_node, "ref_type": ref_type}]
        extract_paths(parent_node, new_path, all_paths)

def classify_path(path):
    rev_path = path[::-1]
    
    main_process = None
    for p in rev_path:
        if p["node"].type in ("process", "processroute"):
            main_process = p["node"]
            break
            
    hierarchy_str = f"[{rev_path[0]['node'].type}] {rev_path[0]['node'].name}"
    
    intermediate_nodes = []
    
    for i in range(1, len(rev_path)):
        prev = rev_path[i-1]
        prev_node = prev["node"]
        curr = rev_path[i]
        curr_node = curr["node"]
        ref_type = prev.get("ref_type", "DEPENDENT")
        
        if ref_type.startswith("SHAPE:"):
            shape_label = ref_type.replace("SHAPE:", "", 1)
            intermediate_nodes.append(shape_label)
        elif ref_type != "DEPENDENT" and ref_type != "CIRCULAR_LINK":
            intermediate_nodes.append(f"[{prev_node.type}] {prev_node.name} -> {ref_type}")
            
        if prev_node.type == "processroute" or ref_type == "INDEPENDENT":
            rel_str = f"-[Process Route]->"
        elif prev_node.type == "process" and curr_node.type == "process":
            rel_str = f"-[Process Call]->"
        elif ref_type.startswith("SHAPE:"):
            shape_label = ref_type.replace("SHAPE:", "", 1)
            if shape_label.startswith("(") and shape_label.endswith(")"):
                rel_str = f"-{shape_label}->"
            else:
                rel_str = f"-({shape_label})->"
        elif ref_type == "(Configuration)":
            rel_str = f"-{ref_type}->"
        else:
            rel_str = f"-({ref_type})->"
            
        hierarchy_str += f" {rel_str} [{curr_node.type}] {curr_node.name}"
        
    intermediate_chain_str = " -> ".join(intermediate_nodes) if intermediate_nodes else "(Directly Referenced)"
        
    return {
        "main_process": main_process,
        "intermediate_chain": intermediate_chain_str,
        "hierarchy_str": hierarchy_str
    }

def format_hierarchy_sheet(ws_det, deduped_rows):
    header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    
    headers = [
        "Main Process", "Main Process ID", "Main Process Deployment",
        "Intermediate Reference Chain", "Full Hierarchy Path"
    ]
    
    for col_idx, h in enumerate(headers, 1):
        cell = ws_det.cell(row=1, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        ws_det.column_dimensions[cell.column_letter].width = 35
        
    ws_det.column_dimensions['E'].width = 110 # Hierarchy Path
    
    deduped_rows.sort(key=lambda x: (x[0], x[3], x[4]))
            
    for row_idx, row_data in enumerate(deduped_rows, 2):
        for col_idx, val in enumerate(row_data, 1):
            ws_det.cell(row=row_idx, column=col_idx, value=val)
            
    if len(deduped_rows) > 1:
        current_mp = None
        start_row = 2
        for i in range(2, len(deduped_rows) + 3):
            mp_val = ws_det.cell(row=i, column=1).value if i <= len(deduped_rows) + 1 else None
            if mp_val != current_mp:
                if current_mp is not None and (i - 1) > start_row:
                    ws_det.merge_cells(start_row=start_row, start_column=1, end_row=i-1, end_column=1)
                    ws_det.merge_cells(start_row=start_row, start_column=2, end_row=i-1, end_column=2)
                    ws_det.merge_cells(start_row=start_row, start_column=3, end_row=i-1, end_column=3)
                    ws_det.cell(row=start_row, column=1).alignment = Alignment(vertical='top')
                    ws_det.cell(row=start_row, column=2).alignment = Alignment(vertical='top')
                    ws_det.cell(row=start_row, column=3).alignment = Alignment(vertical='top')
                current_mp = mp_val
                start_row = i

    ws_det.auto_filter.ref = ws_det.dimensions

def format_summary_sheet(ws_sum, target_node, target_env_name, is_build, deduped_rows, is_deployed_in_env):
    header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    
    unique_main_processes = set()
    for row in deduped_rows:
        mp_id = row[1]
        if mp_id:
            unique_main_processes.add(mp_id)

    if is_build:
        deployment_status = "N/A (Build Report)"
    else:
        deployment_status = "DEPLOYED (Active)" if is_deployed_in_env else f"NOT DEPLOYED in {target_env_name}"

    summary_data = [
        ("Audited Component Name", target_node.name),
        ("Component Type", target_node.type),
        ("Report Scope", "Build (All paths)" if is_build else f"Environment: {target_env_name}"),
        ("Total Main Processes Referenced", len(unique_main_processes)),
        ("Audit Timestamp", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Direct Deployment Status", deployment_status)
    ]
    
    for row_idx, (k, v) in enumerate(summary_data, 1):
        cell_k = ws_sum.cell(row=row_idx, column=1, value=k)
        cell_k.fill = header_fill
        cell_k.font = header_font
        ws_sum.cell(row=row_idx, column=2, value=str(v))
        
    ws_sum.column_dimensions['A'].width = 35
    ws_sum.column_dimensions['B'].width = 50

def generate_build_excel(target_node, all_paths, output_file):
    wb = Workbook()
    ws_sum = wb.active
    ws_sum.title = "Audit Summary"
    
    ws_det = wb.create_sheet(title="Hierarchy Details")
    
    seen_paths = set()
    deduped_rows = []
    
    for path in all_paths:
        classified = classify_path(path)
        main_proc = classified["main_process"]
        mp_name = main_proc.name if main_proc else "None / Orphaned"
        mp_id = main_proc.comp_id if main_proc else ""
        mp_dep = "N/A (Build Scope)"
        
        intermediate_chain = classified["intermediate_chain"]
        hierarchy = classified["hierarchy_str"]
        
        row_tuple = (mp_name, mp_id, mp_dep, intermediate_chain, hierarchy)
        if row_tuple not in seen_paths:
            seen_paths.add(row_tuple)
            deduped_rows.append(row_tuple)
            
    format_hierarchy_sheet(ws_det, deduped_rows)
    format_summary_sheet(ws_sum, target_node, "N/A", True, deduped_rows, False)
    
    wb.save(output_file)
    print(f"Build Audit report generated: {output_file}")

def generate_env_excel(target_node, all_paths, environments, account_id, api_url, username, token, output_file):
    wb = Workbook()
    
    # Remove default sheet
    wb.remove(wb.active)
    
    for env in environments:
        env_id = env["id"]
        env_name = env["name"]
        
        # Safe sheet names
        safe_env_name = "".join([c if c.isalnum() else "_" for c in env_name])[:20]
        
        deployed_pkgs = get_deployed_packages(env_id, account_id, api_url, username, token)
        
        ws_sum = wb.create_sheet(title=f"{safe_env_name}_Summary")
        ws_det = wb.create_sheet(title=f"{safe_env_name}_Details")
        
        seen_paths = set()
        deduped_rows = []
        
        for path in all_paths:
            classified = classify_path(path)
            main_proc = classified["main_process"]
            mp_name = main_proc.name if main_proc else "None / Orphaned"
            mp_id = main_proc.comp_id if main_proc else ""
            
            mp_dep = "N/A"
            if main_proc and main_proc.comp_id in deployed_pkgs:
                d = deployed_pkgs[main_proc.comp_id]
                mp_dep = f"v{d['version']} (Pkg: {d['packageId']})"
            elif main_proc:
                mp_dep = "NOT DEPLOYED"
                
            # Filter out undeployed processes
            if mp_dep == "NOT DEPLOYED" and mp_name != "None / Orphaned":
                continue
    
            intermediate_chain = classified["intermediate_chain"]
            hierarchy = classified["hierarchy_str"]
            
            row_tuple = (mp_name, mp_id, mp_dep, intermediate_chain, hierarchy)
            if row_tuple not in seen_paths:
                seen_paths.add(row_tuple)
                deduped_rows.append(row_tuple)
                
        is_deployed = target_node.comp_id in deployed_pkgs
                
        format_hierarchy_sheet(ws_det, deduped_rows)
        format_summary_sheet(ws_sum, target_node, env_name, False, deduped_rows, is_deployed)
        
    try:
        wb.save(output_file)
        if os.path.exists(output_file):
            print(f"Environment Audit report generated and verified on disk: {output_file}")
        else:
            print(f"ERROR: wb.save completed but file not found on disk: {output_file}")
    except Exception as e:
        print(f"CRITICAL ERROR saving Environment Workbook: {e}")

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Generate Boomi Component Excel Audit Report")
    parser.add_argument("--component-id", required=True, help="Target Component ID to audit")
    parser.add_argument("--environment-id", help="Target Environment ID (Optional).")
    parser.add_argument("--mode", required=True, choices=["BUILD", "ENVIRONMENT", "ALL_ENVIRONMENTS", "BUILD_AND_ENVIRONMENT", "BUILD_AND_ALL_ENVIRONMENTS"], help="Reporting mode.")
    parser.add_argument("--account-id", default=os.getenv("BOOMI_ACCOUNT_ID"), help="Boomi Account ID")
    parser.add_argument("--api-url", default=os.getenv("BOOMI_API_URL", "https://api.boomi.com"), help="Boomi Platform API URL")
    parser.add_argument("--username", default=os.getenv("BOOMI_USERNAME"), help="Boomi Username/Email")
    parser.add_argument("--token", default=os.getenv("BOOMI_API_TOKEN"), help="Boomi API Token")

    args = parser.parse_args()

    if not all([args.account_id, args.username, args.token]):
        sys.stderr.write("ERROR: Missing required credentials (account-id, username, token).\n")
        sys.exit(1)
        
    skill_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    out_dir = os.path.join(skill_dir, "output")
    os.makedirs(out_dir, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")

    print(f"Starting audit for component: {args.component_id}...")
    target_node, node_cache = build_graph(args.component_id, args.account_id, args.api_url, args.username, args.token)
    
    print(f"Graph built. Total unique components found in hierarchy: {len(node_cache)}")
    
    all_paths = []
    extract_paths(target_node, [{"node": target_node}], all_paths)
    print(f"Total hierarchy paths extracted: {len(all_paths)}")
    
    if args.mode in ["BUILD", "BUILD_AND_ENVIRONMENT", "BUILD_AND_ALL_ENVIRONMENTS"]:
        build_out_file = os.path.join(out_dir, f"Build_Audit_Report_{args.component_id}_{timestamp}.xlsx")
        generate_build_excel(target_node, all_paths, build_out_file)
    
    if args.mode in ["ENVIRONMENT", "ALL_ENVIRONMENTS", "BUILD_AND_ENVIRONMENT", "BUILD_AND_ALL_ENVIRONMENTS"]:
        if args.mode in ["ENVIRONMENT", "BUILD_AND_ENVIRONMENT"] and args.environment_id:
            print(f"Generating isolated environment report for Env ID: {args.environment_id}")
            env_out_file = os.path.join(out_dir, f"Environment_Audit_Report_{args.component_id}_{timestamp}.xlsx")
            env_meta = {"id": args.environment_id, "name": f"Env {args.environment_id}"} # Fallback name
            envs = [env_meta]
            # Attempt to get real name
            all_e = get_all_environments(args.account_id, args.api_url, args.username, args.token)
            for e in all_e:
                if e["id"] == args.environment_id:
                    envs = [e]
                    break
            generate_env_excel(target_node, all_paths, envs, args.account_id, args.api_url, args.username, args.token, env_out_file)
        else:
            print("Fetching all available environments...")
            env_out_file = os.path.join(out_dir, f"Environment_Audit_All_Environments_{args.component_id}_{timestamp}.xlsx")
            all_e = get_all_environments(args.account_id, args.api_url, args.username, args.token)
            print(f"Found {len(all_e)} environments.")
            generate_env_excel(target_node, all_paths, all_e, args.account_id, args.api_url, args.username, args.token, env_out_file)
